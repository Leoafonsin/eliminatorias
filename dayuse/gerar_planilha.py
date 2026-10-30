"""Gera DayUse.xlsx com a mesma estrutura que o Code.gs cria no Google Sheets.

Uso:
    python dayuse/gerar_planilha.py                      # planilha vazia
    python dayuse/gerar_planilha.py alunos.xlsx          # já importa os alunos
    python dayuse/gerar_planilha.py alunos.csv

O arquivo de alunos precisa ter colunas com os nomes Nome, CPF, Telefone, Matricula e Email
(maiúsculas/acentos não importam; só Nome é obrigatória).
"""

import csv
import re
import sys
import unicodedata
import uuid
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

SAIDA = Path(__file__).with_name("DayUse.xlsx")

CABECALHOS = {
    "Pessoas": ["ID", "Nome", "Tipo", "CPF", "Telefone", "Email", "Matricula", "Ativo", "DataCadastro", "Observacao"],
    "Presencas": ["ID", "Data", "Turno", "Pessoa", "Nome", "Tipo", "Valor", "Pago", "FormaPagamento",
                  "HoraEntrada", "Observacao"],
    "Precos": ["Chave", "Turno", "Tipo", "Valor"],
}
PRECOS = [("Manhã", "Aluno", 0), ("Manhã", "Visitante", 40), ("Tarde", "Aluno", 30), ("Tarde", "Visitante", 40)]

AZUL = PatternFill("solid", fgColor="1F4E78")
CINZA = PatternFill("solid", fgColor="D9E1F2")
AMARELO = PatternFill("solid", fgColor="FFF2CC")
MOEDA = '"R$" #,##0.00'


def normalizar(texto):
    texto = unicodedata.normalize("NFD", str(texto or "").strip().lower())
    return re.sub(r"\s+", " ", "".join(c for c in texto if unicodedata.category(c) != "Mn"))


def limpar_cpf(cpf):
    if isinstance(cpf, float) and cpf.is_integer():
        cpf = int(cpf)
    digitos = re.sub(r"\D", "", str(cpf or ""))
    return digitos.zfill(11) if 9 <= len(digitos) < 11 else digitos


def limpar_matricula(mat):
    if isinstance(mat, float) and mat.is_integer():
        mat = int(mat)
    return str(mat or "").strip().upper()


def ler_alunos(caminho):
    caminho = Path(caminho)
    if caminho.suffix.lower() == ".csv":
        with open(caminho, encoding="utf-8-sig", newline="") as f:
            amostra = f.read(2048)
            f.seek(0)
            linhas = list(csv.reader(f, delimiter=";" if amostra.count(";") > amostra.count(",") else ","))
    else:
        ws = load_workbook(caminho, data_only=True).active
        linhas = [list(r) for r in ws.iter_rows(values_only=True)]

    cab = [normalizar(c) for c in linhas[0]]

    def col(*nomes):
        return next((cab.index(n) for n in nomes if n in cab), None)

    i_nome = col("nome", "nome completo", "aluno")
    if i_nome is None:
        sys.exit(f"Não achei a coluna Nome. Cabeçalhos encontrados: {linhas[0]}")
    i_cpf, i_tel, i_mat = col("cpf"), col("telefone", "celular", "fone"), col("matricula", "mat")
    i_email = col("email", "e-mail")

    def pegar(linha, i):
        return linha[i] if i is not None and i < len(linha) and linha[i] is not None else ""

    for linha in linhas[1:]:
        yield pegar(linha, i_nome), pegar(linha, i_cpf), pegar(linha, i_tel), pegar(linha, i_mat), pegar(linha, i_email)


def montar_pessoas(alunos):
    cpfs, mats, nomes = set(), set(), set()
    pessoas, avisos, ignorados = [], [], 0
    for nome, cpf, tel, mat, email in alunos:
        nome = str(nome).strip()
        if not nome:
            continue
        cpf, mat, chave = limpar_cpf(cpf), limpar_matricula(mat), normalizar(nome)
        if (cpf and cpf in cpfs) or (mat and mat in mats) or (not cpf and not mat and chave in nomes):
            ignorados += 1
            continue
        if cpf and len(cpf) != 11:
            avisos.append(f"{nome}: CPF com {len(cpf)} dígitos")
        cpfs.add(cpf)
        mats.add(mat)
        nomes.add(chave)
        pessoas.append([uuid.uuid4().hex[:8], nome, "Aluno", cpf, str(tel).strip(), str(email).strip().lower(), mat,
                        True, date.today(), ""])
    return pessoas, avisos, ignorados


def cabecalho(ws, valores, linha=1):
    for c, v in enumerate(valores, 1):
        cel = ws.cell(linha, c, v)
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = AZUL


def lista(ws, intervalo, opcoes):
    dv = DataValidation(type="list", formula1='"' + ",".join(opcoes) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(intervalo)


def gerar(pessoas):
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Pessoas")
    cabecalho(ws, CABECALHOS["Pessoas"])
    for p in pessoas:
        ws.append(p)
    for r in range(2, max(ws.max_row, 2000) + 1):
        for c in (4, 5, 6, 7):
            ws.cell(r, c).number_format = "@"
        ws.cell(r, 9).number_format = "DD/MM/YYYY"
    lista(ws, "C2:C5000", ["Aluno", "Visitante"])
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("Presencas")
    cabecalho(ws, CABECALHOS["Presencas"])
    for r in range(2, 2001):
        ws.cell(r, 2).number_format = "DD/MM/YYYY"
        ws.cell(r, 7).number_format = MOEDA
        ws.cell(r, 10).number_format = "HH:MM"
    lista(ws, "C2:C5000", ["Manhã", "Tarde"])
    lista(ws, "I2:I5000", ["Pix", "Dinheiro", "Cartão", "Isento"])
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("Precos")
    cabecalho(ws, CABECALHOS["Precos"])
    for turno, tipo, valor in PRECOS:
        ws.append([f"{turno}|{tipo}", turno, tipo, valor])
        ws.cell(ws.max_row, 4).number_format = MOEDA

    ws = wb.create_sheet("Resumo")
    ws["A1"] = "RESUMO DO DAY USE"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Data:"
    ws["A2"].font = Font(bold=True)
    ws["B2"] = "=TODAY()"
    ws["B2"].number_format = "DD/MM/YYYY"
    ws["B2"].fill = AMARELO
    ws["C2"] = "← troque a data para ver outro dia"
    ws["C2"].font = Font(color="888888")

    cabecalho(ws, ["Turno", "Tipo", "Pessoas", "Valor previsto", "Valor recebido", "Pendente"], 4)
    for i, (turno, tipo, _) in enumerate(PRECOS):
        r = 5 + i
        filtro = f"Presencas!$B:$B,$B$2,Presencas!$C:$C,$A{r},Presencas!$F:$F,$B{r}"
        ws.append([turno, tipo, f"=COUNTIFS({filtro})", f"=SUMIFS(Presencas!$G:$G,{filtro})",
                   f"=SUMIFS(Presencas!$G:$G,{filtro},Presencas!$H:$H,TRUE)", f"=D{r}-E{r}"])
    ws.append(["TOTAL", "", "=SUM(C5:C8)", "=SUM(D5:D8)", "=SUM(E5:E8)", "=SUM(F5:F8)"])
    for c in range(1, 7):
        ws.cell(9, c).font = Font(bold=True)
        ws.cell(9, c).fill = CINZA
    for r in range(5, 10):
        for c in (4, 5, 6):
            ws.cell(r, c).number_format = MOEDA

    ws["A11"] = "Por forma de pagamento"
    ws["A11"].font = Font(bold=True)
    cabecalho(ws, ["Forma", "Qtd", "Valor"], 12)
    for i, forma in enumerate(["Pix", "Dinheiro", "Cartão", "Isento"]):
        r = 13 + i
        filtro = f"Presencas!$B:$B,$B$2,Presencas!$I:$I,$A{r}"
        ws.cell(r, 1, forma)
        ws.cell(r, 2, f"=COUNTIFS({filtro})")
        ws.cell(r, 3, f"=SUMIFS(Presencas!$G:$G,{filtro})").number_format = MOEDA

    larguras = {"Pessoas": [10, 35, 11, 14, 16, 30, 12, 8, 13, 30],
                "Presencas": [10, 12, 9, 10, 35, 11, 11, 8, 15, 11, 30],
                "Precos": [18, 9, 11, 10], "Resumo": [12, 12, 10, 16, 16, 12]}
    for nome, ls in larguras.items():
        for i, w in enumerate(ls):
            wb[nome].column_dimensions[chr(65 + i)].width = w

    wb.save(SAIDA)


if __name__ == "__main__":
    alunos = ler_alunos(sys.argv[1]) if len(sys.argv) > 1 else []
    pessoas, avisos, ignorados = montar_pessoas(alunos)
    gerar(pessoas)
    print(f"{SAIDA} gerado com {len(pessoas)} aluno(s). {ignorados} repetido(s) ignorado(s).")
    for a in avisos:
        print("Conferir:", a)
